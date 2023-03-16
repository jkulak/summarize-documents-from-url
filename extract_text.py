import magic
import PyPDF2
import pdfplumber
import docx
import openpyxl

def pdf_to_text(file_path):
    plain_text = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            plain_text += page.extract_text()
    return plain_text

def docx_to_text(file_path):
    plain_text = ""
    doc = docx.Document(file_path)
    for paragraph in doc.paragraphs:
        plain_text += paragraph.text + "\n"
    return plain_text

def xlsx_to_text(file_path):
    plain_text = ""
    wb = openpyxl.load_workbook(file_path, read_only=True)
    for sheet in wb:
        for row in sheet.iter_rows():
            plain_text += " ".join([str(cell.value) if cell.value else "" for cell in row]) + "\n"
    return plain_text

def file_to_text(file_path):
    file_type = magic.from_file(file_path, mime=True)

    if file_type == "application/pdf":
        plain_text = pdf_to_text(file_path)
    elif file_type in ["application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/msword"]:
        plain_text = docx_to_text(file_path)
    elif file_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        plain_text = xlsx_to_text(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")

    return plain_text

if __name__ == "__main__":
    file_path = "YOUR_FILE_PATH_HERE"
    text = file_to_text(file_path)